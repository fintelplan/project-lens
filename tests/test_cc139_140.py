"""CC-139 / CC-140 -- order item 11 (L2.4 / L2.5) and a log-hygiene gap.
CC-139: the Regular Report docx carried the model's markdown as characters (Oct 2: "**" x742,
"####" headings, 22 table rows as pipes, 24 "---" rules), a fixed "Mistral-small" subtitle while
ministral-8b-2512 wrote it, and the model's own wrong "Date:" line. WCAG 1.3.1: structure is
programmatically determined -- Word headings, runs, bullets, tables.
CC-140: _redact masked a secret URL's full value only; an exception names host='...' without the
scheme, so the Supabase host reached a public Actions log. Two copies of _redact had drifted.
    python tests/test_cc139_140.py
"""
import os
import sys

ROOT = os.environ.get("LENS_ROOT_DIR") or os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
CODE = os.path.join(ROOT, "code")
sys.path.insert(0, CODE)
PASS = 0
FAIL = 0


def check(name, got, want):
    global PASS, FAIL
    if got == want:
        PASS += 1
        print(f"  PASS  {name}: {got!r}")
    else:
        FAIL += 1
        print(f"  FAIL  {name}: got {got!r} want {want!r}")


import lens_regular_report as R

FIX = "\n".join([
    "**PROJECT LENS REGULAR REPORT**",
    "**Date:** 2026-10-01",
    "**Mission:** *Analyzing* the convergence",
    "---",
    "### **PART 1 \u2014 DETECTION**",
    "#### **1. Injection Patterns Across Domains**",
    "- **FACT_VOID (False Equivalence):** a *quiet* claim",
    "1. **FlyDubai incident** ([REF-1]) triggers fear",
    "| **Actor** | **Legitimacy Gap** |",
    "|---|---|",
    "| US | high |",
    "",
    "Plain closing line.",
])
B = R.md_blocks(FIX)


def texts(b):
    if b[0] == "head":
        yield b[2]
    elif b[0] in ("bullet", "para"):
        for t, _, _ in b[1]:
            yield t
    else:
        for row in b[1]:
            for cell in row:
                for t, _, _ in cell:
                    yield t


allt = [t for b in B for t in texts(b)]
check("no ** left anywhere", sum(t.count("**") for t in allt), 0)
check("no line starts with # or |", [t for t in allt if t.lstrip()[:1] in ("#", "|")], [])
check("the decorative rule is dropped", [t for t in allt if t.strip() == "---"], [])
check("the model's own Date line is dropped", [t for t in allt if t.startswith("Date")], [])
check("#### becomes a heading", ("head", 3, "1. Injection Patterns Across Domains") in B, True)
check("CC-69's PART rule still makes a level-1 heading", ("head", 1, "PART 1 \u2014 DETECTION") in B, True)
check("a bullet keeps its bold and italic as runs",
      ("bullet", [("FACT_VOID (False Equivalence):", True, False), (" a ", False, False),
                  ("quiet", False, True), (" claim", False, False)]) in B, True)
tables = [b for b in B if b[0] == "table"]
check("the pipe rows become one table of 2 rows (separator dropped)",
      [len(t[1]) for t in tables], [2])

src = open(os.path.join(CODE, "lens_regular_report.py"), "rb").read().decode("utf-8")
check("the subtitle names who answered, not a fixed Mistral-small",
      ("Free Tier  |  Mistral-small" in src, "ANSWERED.get('model')" in src), (False, True))
check("render_docx renders the body through md_blocks", "    _render_body(doc, report_text)" in src, True)
try:
    R._log_completion(None, "mistral", "ministral-8b-2512")
except Exception:
    pass
check("the completion log records who answered", R.ANSWERED.get("model"), "ministral-8b-2512")

try:
    from docx import Document
    d = Document()
    R._render_body(d, FIX)
    words = [p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]
    check("in a real docx: no ** and no pipe rows", (sum(w.count("**") for w in words),
          sum(1 for w in words if w.lstrip().startswith("|"))), (0, 0))
except ImportError:
    print("  NOTE  python-docx not installed here: the docx layer is checked where it is")

import lens_provider_refusal as P
import lens_framing_rubrics as F
HOST = "abcdefghijklmnopqrst.supabase.co"
os.environ["SUPABASE_URL"] = "https://" + HOST
os.environ["SHORTHOST_URL"] = "https://ab.co/abcdefghijkl"
os.environ["HARMLESS_SETTING"] = "a_long_but_not_secret_value"
msg = f"ReadTimeout HTTPSConnectionPool(host='{HOST}', port=443)"
check("a URL secret's host is masked when it appears without the scheme", HOST in P._redact(msg), False)
check("framing_rubrics._redact gives the same answer (one implementation)", F._redact(msg), P._redact(msg))
check("a short host is left alone", "ab.co" in P._redact("see ab.co here"), True)
check("a non-secret value is untouched", "a_long_but_not_secret_value" in P._redact("a_long_but_not_secret_value"), True)
check("empty text is safe", P._redact(""), "")
print()
print(f"RESULT: {PASS} passed, {FAIL} failed")
sys.exit(1 if FAIL else 0)
