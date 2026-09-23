"""CC-117 -- the S2 Information Shaping Report is restored, loud and readable.

Sep 21-23 2026: five of five waves, S2-RPT called mistral-small-latest, got 429
on the first call, three times, and returned AI_FAILED to an orchestrator that
threw the result away. No report, no log line saying so, a green step.
    python tests/test_s2_report_loud.py
"""
import os
import sys
import types

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
sys.path.insert(0, os.path.join(ROOT, "code"))

import lens_s2_step_report as R

PASS = 0
FAIL = 0


def check(name, got, want):
    global PASS, FAIL
    if got == want:
        PASS += 1
        print(f"  PASS  {name}: {got!r}")
    else:
        FAIL += 1
        print(f"  FAIL  {name}: got {got!r} want {want!r}")


SRC = open(R.__file__, "rb").read()

print("== the model is not the refused class ==")
check("MODEL", R.MODEL, "ministral-8b-2512")
check("no mistral-small in MODEL", "mistral-small" in R.MODEL, False)
check("headroom: MAX_TOKENS >= 6000 (probe used 3,278)", R.MAX_TOKENS >= 6000, True)
check("finish_reason is logged", b"finish_reason={fr}" in SRC, True)
check('the docx no longer claims "Mistral-small"', b'Mistral-small")' in SRC, False)

missing = [n for n in ("announce_failure", "_plain", "_add_runs") if not hasattr(R, n)]
check("CC-117 functions exist", missing, [])
if missing:
    print()
    print(f"RESULT: {PASS} passed, {FAIL} failed (stopped: the rest needs the functions)")
    sys.exit(1)

print("== a refusal is retried, then said and recorded ==")
os.environ["MISTRAL_API_KEY"] = "test-not-a-key"
posts, sleeps, recorded = [], [], []


class _Resp:
    def __init__(self, code, body):
        self.status_code = code
        self.text = str(body)
        self._body = body

    def json(self):
        return self._body


def _post_429(*a, **k):
    posts.append(k.get("json", {}).get("model"))
    return _Resp(429, {"message": "Service tier capacity exceeded for this model."})


_real_post, _real_sleep = R.requests.post, R.time.sleep
fake_ref = types.ModuleType("lens_provider_refusal")
fake_ref.record_refusal = lambda provider, model, exc=None, status=None, text=None: recorded.append((provider, model, status))
_real_ref = sys.modules.get("lens_provider_refusal")
sys.modules["lens_provider_refusal"] = fake_ref
R.requests.post = _post_429
R.time.sleep = lambda s: sleeps.append(s)
try:
    got = R.call_mistral("prompt")
finally:
    R.requests.post, R.time.sleep = _real_post, _real_sleep
check("429 x3 -> None", got, None)
check("three attempts, all on the new model", posts, ["ministral-8b-2512"] * 3)
check("no sleep after the last attempt", sleeps, [20, 40])
check("one PROVIDER_REFUSAL recorded", recorded, [("mistral", "ministral-8b-2512", 429)])

print("== an answer is returned; an empty one is not ==")
posts.clear(); sleeps.clear(); recorded.clear()
R.requests.post = lambda *a, **k: _Resp(200, {"choices": [{"message": {"content": " PART A \u2014 X "}, "finish_reason": "stop"}]})
R.time.sleep = lambda s: sleeps.append(s)
try:
    check("200 -> text", R.call_mistral("prompt"), "PART A \u2014 X")
    R.requests.post = lambda *a, **k: _Resp(200, {"choices": [{"message": {"content": ""}, "finish_reason": "length"}]})
    check("200 empty -> None", R.call_mistral("prompt"), None)
    check("an empty 200 is not recorded as a provider refusal", recorded, [])
finally:
    R.requests.post, R.time.sleep = _real_post, _real_sleep
    if _real_ref is not None:
        sys.modules["lens_provider_refusal"] = _real_ref
    else:
        sys.modules.pop("lens_provider_refusal", None)

print("== a report that did not arrive is announced ==")
sent = []
_real_tt = R.send_telegram_text
R.send_telegram_text = lambda t: sent.append(t) or True
try:
    check("COMPLETE -> no announcement", R.announce_failure({"status": "COMPLETE"}), False)
    check("COMPLETE sent nothing", sent, [])
    for st in ("AI_FAILED", "NO_DATA", "DOCX_FAILED", "SEND_FAILED", "ERROR"):
        sent.clear()
        check(st + " -> announced", R.announce_failure({"status": st}), True)
        check(st + " message names it", bool(sent) and ("FAILED" in sent[0] and st in sent[0]), True)
    sent.clear()
    check("no status at all -> announced", R.announce_failure({}), True)
finally:
    R.send_telegram_text = _real_tt

print("== the model's markdown headings render as headings ==")
check("## **PART A ...**", R._plain("## **PART A \u2014 THE INJECTION ARCHITECTURE**"),
      "PART A \u2014 THE INJECTION ARCHITECTURE")
check("### **sub**", R._plain("### **The Overall Contamination Picture**"), "The Overall Contamination Picture")
check("bare PART unchanged", R._plain("PART F \u2014 MISSION ANALYST SYNTHESIS"), "PART F \u2014 MISSION ANALYST SYNTHESIS")


class _Para:
    def __init__(self):
        self.runs = []

    def add_run(self, t):
        r = types.SimpleNamespace(text=t, bold=None)
        self.runs.append(r)
        return r


p = _Para()
R._add_runs(p, "1. **US/NATO strength**\u2014by highlighting **deterrence collapse**.")
check("bold runs", [(r.text, bool(r.bold)) for r in p.runs],
      [("1. ", False), ("US/NATO strength", True), ("\u2014by highlighting ", False),
       ("deterrence collapse", True), (".", False)])
try:
    import docx  # noqa: F401
    sample = ("# **INTELLIGENCE REPORT**\n**Run ID: x**\n\n"
              + "\n".join("## **PART %s \u2014 T%s**\nbody **b** text" % (c, c) for c in "ABCDEF"))
    path = R.render_docx(sample, "2026-09-23", {"threat_level": "CRITICAL", "quality_score": 0.5})
    d = docx.Document(path)
    h1 = [x.text for x in d.paragraphs if x.style.name == "Heading 1"]
    check("six PART headings in the docx", len(h1), 6)
    check("no markdown left in the docx", any(("**" in x.text or x.text.startswith("#")) for x in d.paragraphs), False)
    check("subtitle names the model", any("ministral-8b-2512" in x.text for x in d.paragraphs), True)
    os.remove(path)
except ImportError:
    print("  NOTE  python-docx not installed here: the docx render check did not run")

print("== the orchestrator reads the outcome, and red skips nothing ==")
orc = open(os.path.join(ROOT, "code", "lens_s2_orchestrator.py"), "rb").read()
check("orchestrator no longer discards the result", b"from lens_s2_step_report import run_s2_report" in orc, False)
check("orchestrator prints the report status", b"[S2-ORC] S2 Shaping Report: {report_status}" in orc, True)
check("orchestrator exits red on a missing report",
      b'if report_status != "COMPLETE":\r\n        # CC-117' in orc.replace(b"\r\n", b"\n").replace(b"\n", b"\r\n"), True)
wf = open(os.path.join(ROOT, ".github", "workflows", "lens-manage-analyze.yml"), "rb").read()
check("premise: S2, S3 and health steps all run under !cancelled() (red skips nothing)",
      wf.count(b"if: ${{ !cancelled() }}") >= 3, True)

print()
print(f"RESULT: {PASS} passed, {FAIL} failed")
sys.exit(1 if FAIL else 0)
